"""Extract plain text from uploaded PDF, DOCX, and TXT files."""

import io
import re

from docx import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class ExtractionError(ValueError):
    pass


def _extension(filename: str) -> str:
    dot = filename.rfind(".")
    return filename[dot:].lower() if dot != -1 else ""


def _from_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ExtractionError("This PDF is password-protected.") from exc
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _from_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    # Many CV templates put content in tables (two-column layouts).
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _from_txt(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def clean_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def extract_text(uploaded_file) -> str:
    """Return cleaned text from a Django UploadedFile."""
    ext = _extension(uploaded_file.name)
    if ext not in SUPPORTED_EXTENSIONS:
        raise ExtractionError(
            f"Unsupported file type '{ext or 'unknown'}'. Use PDF, DOCX, or TXT."
        )
    data = uploaded_file.read()
    try:
        if ext == ".pdf":
            text = _from_pdf(data)
        elif ext == ".docx":
            text = _from_docx(data)
        else:
            text = _from_txt(data)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(
            f"Could not read '{uploaded_file.name}'. The file may be corrupt."
        ) from exc

    text = clean_text(text)
    if len(text) < 50:
        raise ExtractionError(
            f"Very little text found in '{uploaded_file.name}'. "
            "If it is a scanned PDF, export it with selectable text first."
        )
    return text
